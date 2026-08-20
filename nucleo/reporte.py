"""
Generación del reporte Excel de la conciliación.

El reporte es lo que queda como respaldo del cierre, así que está pensado
para que se entienda sin tener el sistema al lado: la primera hoja muestra
el cálculo completo paso a paso, y las siguientes el detalle que respalda
cada número.

Hojas:
    1. Conciliación  — el cálculo: saldo del libro, apertura, ajustes,
                       saldo del banco y diferencia final.
    2. Pendientes    — lo que no se pudo emparejar, con notas.
    3. Conciliados   — los movimientos emparejados, para auditar.

Reemplaza a `generar_reporte` de conciliacion.py, que mostraba etiquetas
que no correspondían con lo que realmente calculaba el sistema (por
ejemplo, rotulaba el saldo acumulado como "movimientos netos del mes").
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# --- Estilos ---------------------------------------------------------
TINTA = "1F2933"
ACENTO = "1B4965"
GRIS_SUAVE = "EEF1F5"
VERDE = "1F6B45"
AMBAR = "8A5510"

PESOS = '"$" #,##0.00;[Red]-"$" #,##0.00'

F_TITULO = Font(name="Calibri", size=15, bold=True, color=TINTA)
F_SUB = Font(name="Calibri", size=10, color="6B7885")
F_SECCION = Font(name="Calibri", size=11, bold=True, color=ACENTO)
F_ENCABEZADO = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name="Calibri", size=10, color=TINTA)
F_FUERTE = Font(name="Calibri", size=11, bold=True, color=TINTA)
F_TOTAL = Font(name="Calibri", size=12, bold=True, color=ACENTO)

R_ENCABEZADO = PatternFill("solid", fgColor=ACENTO)
R_SUAVE = PatternFill("solid", fgColor=GRIS_SUAVE)

_linea = Side(style="thin", color="D0D7DE")
BORDE_INF = Border(bottom=_linea)
BORDE_TOTAL = Border(top=Side(style="medium", color=TINTA))


def generar_reporte_excel(crm, banco, matches, discrepancias, stats, salida,
                          nombre_banco="", periodo=""):
    """
    Escribe el reporte completo en `salida` (ruta o BytesIO).

    `stats` es el dict de estadísticas que devuelve el orquestador.
    """
    wb = Workbook()
    _hoja_conciliacion(wb.active, stats, nombre_banco, periodo)
    _hoja_pendientes(wb.create_sheet("Pendientes"), discrepancias)
    _hoja_conciliados(wb.create_sheet("Conciliados"), crm, banco, matches)
    wb.save(salida)
    return salida


# ---------------------------------------------------------------------
# Hoja 1 — el cálculo
# ---------------------------------------------------------------------
def _hoja_conciliacion(ws, stats, nombre_banco, periodo):
    ws.title = "Conciliación"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 60

    f = 2
    ws.cell(f, 2, f"Conciliación bancaria{' — ' + nombre_banco if nombre_banco else ''}").font = F_TITULO
    f += 1
    if periodo:
        ws.cell(f, 2, f"Período: {periodo}").font = F_SUB
        f += 1

    dif = stats.get("diferencia_final", 0)
    cierra = abs(dif) <= 1500
    f += 1
    celda = ws.cell(f, 2, "La conciliación cierra" if cierra else "Queda una diferencia por explicar")
    celda.font = Font(name="Calibri", size=11, bold=True, color=VERDE if cierra else AMBAR)
    ws.cell(f, 3, dif).font = Font(name="Calibri", size=11, bold=True, color=VERDE if cierra else AMBAR)
    ws.cell(f, 3).number_format = PESOS
    f += 2

    # --- El cálculo ---
    ws.cell(f, 2, "CÓMO SE LLEGA A LA DIFERENCIA").font = F_SECCION
    f += 1
    for col, txt in ((2, "Concepto"), (3, "Monto"), (4, "De dónde sale")):
        c = ws.cell(f, col, txt)
        c.font = F_ENCABEZADO
        c.fill = R_ENCABEZADO
        c.alignment = Alignment(horizontal="right" if col == 3 else "left")
    f += 1

    origen = stats.get("apertura_origen", "manual")
    if origen == "derivada":
        nota_apertura = (
            f"Derivada: extracto al cierre anterior "
            f"({_num(stats.get('saldo_extracto_anterior'))}) menos el saldo con "
            f"que abre el libro ({_num(stats.get('saldo_crm_arranque'))})"
        )
    else:
        nota_apertura = "Ingresada manualmente"

    filas = [
        ("Saldo según el libro mayor", stats.get("saldo_crm", 0), "Último acumulado del mayor exportado de GBP", False),
        ("+ Saldo de apertura", stats.get("saldo_apertura", 0), nota_apertura, False),
    ]
    for aj in stats.get("ajustes_manuales", []):
        filas.append((f"    {aj.get('concepto', '')}", aj.get("monto", 0),
                      f"{aj.get('cantidad_mov', 0)} movimiento(s)" if aj.get("cantidad_mov") else "Ajuste cargado en la conciliación",
                      False))
    filas.append(("Saldo bancario calculado", stats.get("saldo_banco_calculado", 0),
                  "Lo que debería mostrar el banco según la contabilidad", True))
    filas.append(("Saldo según el extracto", stats.get("saldo_extracto", 0),
                  "Saldo de cierre del extracto bancario", False))

    for etiqueta, monto, nota, destacar in filas:
        cb = ws.cell(f, 2, etiqueta)
        cc = ws.cell(f, 3, monto)
        cd = ws.cell(f, 4, nota)
        cb.font = F_FUERTE if destacar else F_NORMAL
        cc.font = F_FUERTE if destacar else F_NORMAL
        cc.number_format = PESOS
        cd.font = F_SUB
        cd.alignment = Alignment(wrap_text=True, vertical="center")
        if destacar:
            for col in (2, 3, 4):
                ws.cell(f, col).fill = R_SUAVE
        for col in (2, 3, 4):
            ws.cell(f, col).border = BORDE_INF
        f += 1

    cb = ws.cell(f, 2, "DIFERENCIA FINAL")
    cc = ws.cell(f, 3, dif)
    cb.font = F_TOTAL
    cc.font = F_TOTAL
    cc.number_format = PESOS
    ws.cell(f, 4, "Tolerancia aceptada: hasta $1.500,00").font = F_SUB
    for col in (2, 3, 4):
        ws.cell(f, col).border = BORDE_TOTAL
    f += 2

    # --- Resumen del emparejamiento ---
    ws.cell(f, 2, "RESUMEN DEL EMPAREJAMIENTO").font = F_SECCION
    f += 1
    for etiqueta, valor in [
        ("Movimientos en el libro mayor", stats.get("total_crm", 0)),
        ("Movimientos en el extracto", stats.get("total_banco", 0)),
        ("Emparejados", stats.get("total_matches", 0)),
        ("Sin emparejar (ver hoja Pendientes)", stats.get("total_discrepancias", 0)),
    ]:
        ws.cell(f, 2, etiqueta).font = F_NORMAL
        c = ws.cell(f, 3, valor)
        c.font = F_NORMAL
        c.number_format = "#,##0"
        f += 1


# ---------------------------------------------------------------------
# Hoja 2 — lo que quedó pendiente
# ---------------------------------------------------------------------
def _hoja_pendientes(ws, discrepancias):
    ws.sheet_view.showGridLines = False
    encabezados = ["Origen", "Fecha", "Monto", "Descripción", "Contraparte", "Clasificación", "Nota"]
    anchos = [10, 12, 18, 46, 30, 32, 52]
    _escribir_encabezados(ws, encabezados, anchos)

    ordenadas = sorted(discrepancias, key=lambda d: -abs(d.get("monto", 0)))
    f = 2
    for d in ordenadas:
        ws.cell(f, 1, d.get("origen", "")).font = F_NORMAL
        c = ws.cell(f, 2, d.get("fecha"))
        c.number_format = "DD/MM/YYYY"
        c.font = F_NORMAL
        cm = ws.cell(f, 3, d.get("monto", 0))
        cm.number_format = PESOS
        cm.font = F_NORMAL
        ws.cell(f, 4, d.get("descripcion", "")).font = F_NORMAL
        ws.cell(f, 5, d.get("contraparte", "")).font = F_NORMAL
        ws.cell(f, 6, d.get("tipo", "")).font = F_NORMAL
        nota = ws.cell(f, 7, d.get("nota", ""))
        nota.font = Font(name="Calibri", size=9, color=AMBAR if d.get("cruza_meses") else "6B7885")
        nota.alignment = Alignment(wrap_text=True, vertical="center")
        f += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{max(f - 1, 1)}"


# ---------------------------------------------------------------------
# Hoja 3 — lo emparejado
# ---------------------------------------------------------------------
def _hoja_conciliados(ws, crm, banco, matches):
    ws.sheet_view.showGridLines = False
    encabezados = ["ID", "Tipo de coincidencia", "Confianza", "Fecha libro", "Monto libro",
                   "Fecha banco", "Monto banco", "Detalle del banco"]
    anchos = [10, 40, 12, 13, 18, 13, 18, 46]
    _escribir_encabezados(ws, encabezados, anchos)

    f = 2
    for m in matches:
        i_crm, i_bco = m.get("i_crm"), m.get("i_bco")
        ws.cell(f, 1, m.get("match_id", "")).font = F_NORMAL
        ws.cell(f, 2, m.get("tipo", "")).font = F_NORMAL
        ws.cell(f, 3, m.get("confianza", "")).font = F_NORMAL
        if i_crm is not None and i_crm in crm.index:
            c = ws.cell(f, 4, crm.at[i_crm, "fecha"]); c.number_format = "DD/MM/YYYY"; c.font = F_NORMAL
            c = ws.cell(f, 5, float(crm.at[i_crm, "monto"])); c.number_format = PESOS; c.font = F_NORMAL
        if i_bco is not None and i_bco in banco.index:
            c = ws.cell(f, 6, banco.at[i_bco, "fecha"]); c.number_format = "DD/MM/YYYY"; c.font = F_NORMAL
            c = ws.cell(f, 7, float(banco.at[i_bco, "monto"])); c.number_format = PESOS; c.font = F_NORMAL
            ws.cell(f, 8, str(banco.at[i_bco, "descripcion_orig"])[:120]).font = F_NORMAL
        f += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(f - 1, 1)}"


# ---------------------------------------------------------------------
def _escribir_encabezados(ws, encabezados, anchos):
    for i, (txt, ancho) in enumerate(zip(encabezados, anchos), start=1):
        c = ws.cell(1, i, txt)
        c.font = F_ENCABEZADO
        c.fill = R_ENCABEZADO
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 22


def _num(valor):
    """Formatea un número para las notas explicativas."""
    if valor is None:
        return "—"
    return f"${valor:,.2f}"
